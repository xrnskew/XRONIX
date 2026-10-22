"""Платёжеспособность региона — Росстат: среднедушевые денежные доходы населения (руб. в месяц, последний полный
год) и оборот розничной торговли на душу населения (руб. в год). Дополнение к collect_rosstat.py.

    python -m knowledge.collect_rosstat_extra              # из папки ai: скачать то, чего нет в raw/, и разобрать
    python -m knowledge.collect_rosstat_extra --offline    # только из knowledge/raw/rosstat

Результат — knowledge/raw/rosstat/income.json и retail.json: {регион: {"value", "year", "title"}}.
"""
import argparse
import datetime as dt
import re
import sys

from .collect_rosstat import WITHOUT_AO
from .common import RAW_DIR, http, load_json, region_key, save_json

BASE = "https://rosstat.gov.ru/storage/mediabank/"
SOURCES = {
    "income": {"file": "urov_10subg_1pg2026.xlsx", "page": "https://rosstat.gov.ru/folder/13397",
               "title": "Среднедушевые денежные доходы населения по субъектам Российской Федерации, руб./месяц"},
    "retail": {"file": "Oborot_dusha_2024.xlsx", "page": "https://rosstat.gov.ru/statistics/roznichnayatorgovlya",
               "title": "Оборот розничной торговли в расчете на душу населения по субъектам Российской Федерации, рублей"},
}


def download(name: str, offline: bool):
    path = RAW_DIR / "rosstat" / SOURCES[name]["file"]
    if not path.exists():
        if offline:
            sys.exit(f"нет {path}, запустите без --offline")
        path.write_bytes(http(BASE + SOURCES[name]["file"], russian_ca=True, timeout=120))
    return path


def _year(cell):
    m = re.match(r"\s*(\d{4})", str(cell)) if cell is not None else None
    return int(m.group(1)) if m else None


def _regions(rows, col, year) -> dict:
    out = {}
    for r in rows:
        name, value = r[0], r[col] if len(r) > col else None
        if isinstance(value, str):  # «290205²⁾» — число со сноской
            digits = re.match(r"\s*([\d\s.,]+)", value)
            value = float(digits.group(1).replace(" ", "").replace(",", ".")) if digits else None
        if isinstance(name, str) and isinstance(value, (int, float)) and "федеральный округ" not in name.lower() \
                and "российская федерация" not in name.lower():
            out[region_key(name.replace("в том числе", ""))] = {"value": round(value), "year": year, "title": name.strip()}
    return out


def parse_income(path) -> dict:
    """Шапка: строка с годами (год стоит в первой колонке группы) и строка периодов («год», «I кв.»…).
    Берём колонку «год» последнего года, где она есть."""
    import openpyxl
    rows = list(openpyxl.load_workbook(path, read_only=True, data_only=True).worksheets[-1].iter_rows(values_only=True))
    hi = next(i for i, r in enumerate(rows) if any(_year(c) for c in r[1:]))
    years, cur = [], None
    for c in rows[hi]:
        cur = _year(c) or cur
        years.append(cur)
    periods = rows[hi + 1]
    cols = [i for i, p in enumerate(periods) if isinstance(p, str) and p.strip().lower() == "год" and years[i]]
    col = max(cols, key=lambda i: years[i])
    return _regions(rows[hi + 2:], col, years[col])


def parse_retail(path) -> dict:
    import openpyxl
    ws = openpyxl.load_workbook(path, read_only=True, data_only=True)["2"]
    rows = list(ws.iter_rows(values_only=True))
    hi = next(i for i, r in enumerate(rows) if sum(1 for c in r[1:] if _year(c)) > 5)
    col = max((i for i, c in enumerate(rows[hi]) if _year(c)), key=lambda i: _year(rows[hi][i]))
    return _regions(rows[hi + 1:], col, _year(rows[hi][col]))


def value_for(region: str, table: dict):
    key = region_key(region)
    return table.get(WITHOUT_AO.get(key, key)) or table.get(key)


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("--offline", action="store_true")
    args = ap.parse_args(argv)
    for name, parse in (("income", parse_income), ("retail", parse_retail)):
        path = download(name, args.offline)
        regions = parse(path)
        out = RAW_DIR / "rosstat" / f"{name}.json"
        prev = load_json(out, {})
        retrieved = prev.get("retrieved_at") if prev.get("file") == SOURCES[name]["file"] else None
        save_json(out, {"source": BASE + SOURCES[name]["file"], "page": SOURCES[name]["page"], "file": SOURCES[name]["file"],
                        "title": SOURCES[name]["title"], "retrieved_at": retrieved or dt.date.today().isoformat(),
                        "regions": regions})
        year = next(iter(regions.values()))["year"]
        print(f"Росстат, {name}: {len(regions)} регионов, год {year} -> {out.relative_to(RAW_DIR.parent)}")


if __name__ == "__main__":
    main()
